# -*- coding: utf-8 -*-
"""
生成虚构资方的脱敏样例文档（红线：全部为虚构公司，无任何真实机构信息）。

产出 3 份 docx，模拟真实资方对接时的文档包结构：
1. 华信银行_对公收付接口文档     —— 典型银行代收/代付接口（鉴权/接口清单/字段表）
2. 华信银行_日末文件格式说明     —— 日末对账文件格式（含"逾期文件"列定义，验收问题在此）
3. 瑞金租赁_资金对接接口文档     —— 另一家资方，字段命名习惯不同（阶段4映射抽取的素材）

注意：文档标题一律用 Word 内置 Heading 1/2/3 样式——
后端 DocxDocumentParserStrategy 依赖标题样式做"按标题段落"切分。

用法: E:/install/py-env/Scripts/python.exe gen_sample_docs.py
"""
import os
from docx import Document

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_docs")
os.makedirs(OUT_DIR, exist_ok=True)


def add_kv_table(doc, headers, rows):
    """带表头的网格表格。表头行 = 字段说明的列名。"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for c, h in enumerate(headers):
        table.rows[0].cells[c].text = h
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row):
            table.rows[r].cells[c].text = str(v)
    return table


def gen_huaxin_api():
    doc = Document()
    doc.add_heading("华信银行 对公收付接口文档", level=1)

    doc.add_heading("1 接入约定", level=2)
    doc.add_paragraph("本接口用于合作机构通过 HTTP 通道发起对公账户的代收与代付交易。生产环境网关地址为 https://fake-huaxin.example.com/gateway，联调环境为 https://fake-huaxin-uat.example.com/gateway。请求报文统一使用 UTF-8 编码的 JSON。")
    doc.add_paragraph("所有接口采用 HTTPS 双向证书认证。合作机构需向我行提供机构证书公钥，我行分配合作机构编号（merNo）与接入密钥（secretKey）。")

    doc.add_heading("1.1 通用请求头", level=3)
    add_kv_table(doc, ["参数名", "类型", "必填", "说明"], [
        ["Content-Type", "string", "是", "固定值 application/json"],
        ["X-HX-Timestamp", "string", "是", "请求时间戳，格式 yyyyMMddHHmmss，与我行服务器时间偏差不得超过 5 分钟"],
        ["X-HX-Sign", "string", "是", "签名，规则见 1.2 节"],
        ["X-HX-MerNo", "string", "是", "合作机构编号"],
    ])

    doc.add_heading("1.2 签名规则", level=3)
    doc.add_paragraph("将请求体 JSON 按 key 的字典序升序排列后拼接为 key1=value1&key2=value2 形式，末尾拼接 secretKey，对整个字符串做 SHA-256 摘要，十六进制大写即为 X-HX-Sign。签名失败返回错误码 SIGN_ERROR。")

    doc.add_heading("2 单笔代收接口", level=2)
    doc.add_paragraph("代收指从我行企业客户的对公账户扣划资金至合作机构归集账户。单笔限额 500 万元，日累计限额 5000 万元。接口路径 /v1/collections/single。")
    add_kv_table(doc, ["参数名", "类型", "必填", "说明"], [
        ["outBizNo", "string(32)", "是", "合作机构侧业务流水号，需保证唯一"],
        ["payerAccount", "string(32)", "是", "付款方对公账户号"],
        ["payerName", "string(128)", "是", "付款方户名"],
        ["payeeAccount", "string(32)", "是", "收款归集账户号"],
        ["amount", "decimal(16,2)", "是", "交易金额，单位元，保留两位小数"],
        ["remark", "string(200)", "否", "附言"],
    ])
    doc.add_paragraph("响应中的核心字段：hxSeqNo 为我行流水号，status 取值 PROCESSING/SUCCESS/FAIL。最终态以对账文件或交易查询接口为准，状态为 PROCESSING 时不得直接作为失败处理。")

    doc.add_heading("3 单笔代付接口", level=2)
    doc.add_paragraph("代付指从我行归集账户付款至指定对公账户。接口路径 /v1/payments/single。请求字段与代收接口一致，另有 feeFlag 字段：手续费承担方，取值 PAYER/PAYEE，默认 PAYER。")

    doc.add_heading("4 交易查询接口", level=2)
    doc.add_paragraph("接口路径 /v1/transactions/query。支持按 outBizNo 或 hxSeqNo 查询，二者至少传一个。查询仅返回近 90 天内的交易。")
    add_kv_table(doc, ["响应字段", "类型", "说明"], [
        ["hxSeqNo", "string(32)", "我行流水号"],
        ["outBizNo", "string(32)", "机构侧流水号"],
        ["status", "string", "PROCESSING/SUCCESS/FAIL"],
        ["failCode", "string", "失败码，status=FAIL 时返回"],
        ["failMsg", "string(256)", "失败原因描述"],
        ["completeTime", "string(14)", "完成时间 yyyyMMddHHmmss"],
    ])

    doc.add_heading("5 错误码表", level=2)
    add_kv_table(doc, ["错误码", "含义", "处理建议"], [
        ["SIGN_ERROR", "验签失败", "检查签名规则与 secretKey"],
        ["DUPLICATE_OUTBIZNO", "机构流水号重复", "查询原交易状态，勿重复发起"],
        ["AMOUNT_LIMIT", "超过限额", "单笔/日累计限额见 2 节"],
        ["ACCOUNT_NOT_EXIST", "账户不存在", "核对账户号与户名"],
    ])
    path = os.path.join(OUT_DIR, "华信银行_对公收付接口文档_v1.0.docx")
    doc.save(path)
    return path


def gen_huaxin_eod():
    doc = Document()
    doc.add_heading("华信银行 日末文件格式说明", level=1)

    doc.add_heading("1 文件通用约定", level=2)
    doc.add_paragraph("日末文件为每日 23:30 生成，通过 SFTP 推送至机构目录，文件编码 GBK，分隔符为竖线 |，首行为列名行，不包含在数据行内。文件命名规则：文件类型缩写_机构编号_yyyyMMdd.txt，例如 OVERDUE_HX0001_20260930.txt。")
    doc.add_paragraph("金额字段一律为字符串型数字，保留两位小数，不带千分位。日期字段格式 yyyyMMdd。空字段以两个连续分隔符表示。")

    doc.add_heading("2 还款明细文件", level=2)
    doc.add_paragraph("文件类型缩写 REPAY。包含当日所有成功的还款流水，字段定义如下。")
    add_kv_table(doc, ["列序号", "字段名", "类型", "说明"], [
        [1, "还款流水号", "string(32)", "当日唯一"],
        [2, "借据编号", "string(32)", "还款对应的借据"],
        [3, "还款日期", "date", "格式 yyyyMMdd"],
        [4, "还款本金", "decimal", "两位小数"],
        [5, "还款利息", "decimal", "两位小数"],
        [6, "还款渠道", "string", "BIND/ACTIVE/OFFLINE"],
    ])

    doc.add_heading("3 逾期文件", level=2)
    doc.add_paragraph("文件类型缩写 OVERDUE。每日日末生成，包含全部处于逾期状态的借据快照。机构收到文件后应与本地逾期台账比对，差异以我行文件为准。字段定义如下。")
    add_kv_table(doc, ["列序号", "字段名", "类型", "说明"], [
        [1, "借据编号", "string(32)", "逾期借据的唯一编号"],
        [2, "起租逾期日", "date", "首期逾期发生日，格式 yyyyMMdd"],
        [3, "逾期天数", "int", "截至文件生成日的连续逾期天数，含当日"],
        [4, "逾期本金", "decimal", "逾期未还本金余额，两位小数"],
        [5, "逾期利息", "decimal", "逾期未还利息余额，两位小数"],
        [6, "逾期阶段", "string", "M1/M2/M3/M4PLUS，按逾期天数划分"],
    ])
    doc.add_paragraph("逾期阶段划分规则：M1 为逾期 1-30 天，M2 为 31-60 天，M3 为 61-90 天，M4PLUS 为 91 天及以上。逾期阶段与催收策略挂钩，机构侧不得自行修改。")

    doc.add_heading("4 差错处理", level=2)
    doc.add_paragraph("机构发现文件金额与本地台账不一致时，应在 T+1 日 18:00 前通过差错登记接口反馈，我行在 T+2 日回复核查结论。连续 3 日未反馈视为确认。")
    path = os.path.join(OUT_DIR, "华信银行_日末文件格式说明_v1.0.docx")
    doc.save(path)
    return path


def gen_ruijin_api():
    doc = Document()
    doc.add_heading("瑞金租赁 资金对接接口文档", level=1)

    doc.add_heading("一、接口总则", level=2)
    doc.add_paragraph("瑞金租赁（虚构）资金对接采用 REST 风格，网关地址 https://fake-ruijin.example.com/openapi。与华信银行不同，本家资方要求所有请求放置于 JSON 包体中，使用 RSA2 私钥签名，签名放置于请求体字段 sign 中而非请求头。")

    doc.add_heading("（一）商户准入", level=3)
    doc.add_paragraph("合作机构接入前需完成商户备案，取得 partner_id 与 partner_key。partner_key 仅用于本地生成签名，不得随报文传输。")

    doc.add_heading("二、放款指令接口", level=2)
    doc.add_paragraph("接口路径 /openapi/loan/order/create。用于合作机构向瑞金租赁发起单笔放款指令。指令受理成功返回受理编号，放款结果异步通知。")
    add_kv_table(doc, ["参数名称", "类型", "是否必填", "备注"], [
        ["out_order_no", "string(64)", "是", "机构放款订单号"],
        ["contract_no", "string(64)", "是", "融资租赁合同号"],
        ["loan_amount", "string(20)", "是", "放款金额，字符串型数字，单位元"],
        ["bank_account_no", "string(48)", "是", "收款账户"],
        ["bank_name", "string(128)", "是", "收款银行全称"],
        ["purpose", "string(200)", "是", "资金用途，按合同填写"],
    ])
    doc.add_paragraph("异步通知字段：notify_type 取值 LOAN_RESULT；loan_status 取值 SUCCESS / FAILED / PROCESSING；fail_reason 在失败时返回中文描述。通知需在 24 小时内应答 SUCCESS，否则触发重试，最多重试 5 次。")

    doc.add_heading("三、对账单文件", level=2)
    doc.add_paragraph("瑞金租赁不提供日末 SFTP 文件，改为 T+1 日上午 10 点前开放对账单下载接口 /openapi/reconcile/download，参数 biz_date 指定对账日期，返回 CSV 内容，UTF-8 编码，逗号分隔。")
    add_kv_table(doc, ["字段位置", "字段名称", "类型", "说明"], [
        ["第1列", "order_no", "string", "放款订单号"],
        ["第2列", "settle_date", "date", "清算日期 yyyy-MM-dd"],
        ["第3列", "settle_amount", "string", "实际清算金额"],
        ["第4列", "fee_amount", "string", "手续费金额"],
    ])
    path = os.path.join(OUT_DIR, "瑞金租赁_资金对接接口文档_v1.0.docx")
    doc.save(path)
    return path


if __name__ == "__main__":
    for p in [gen_huaxin_api(), gen_huaxin_eod(), gen_ruijin_api()]:
        print("生成:", p)
